"""Generate a synthetic MRP workbook; no business data or model calls are used."""

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def main() -> None:
    output = Path(__file__).resolve().with_name("sample_mrp.xlsx")
    workbook = Workbook()
    workbook.properties.creator = "MRP Insight Console"
    workbook.properties.title = "Synthetic MRP demo data"
    sheet = workbook.active
    sheet.title = "MRP Demo"
    sheet.append([
        "Material Number",
        "plant_id",
        "mrp_element_receipt_or_requirement_date",
        "Mrp Element Short Name",
        "Qty",
    ])

    start = datetime(2026, 1, 5)
    for material, initial_stock in [
        ("DEMO-STABLE", 500),
        ("DEMO-SHORTAGE", 220),
        ("DEMO-SPIKE", 350),
    ]:
        sheet.append([material, "DEMO-PLANT", start, "Stock", initial_stock])
        for week in range(32):
            demand = 75 + (week % 4) * 5
            supply = demand
            if material == "DEMO-SHORTAGE":
                supply = 45
            elif material == "DEMO-SPIKE":
                if week in (12, 13):
                    demand += 320
                if week == 18:
                    supply += 650
            monday = start + timedelta(weeks=week)
            sheet.append([material, "DEMO-PLANT", monday + timedelta(days=1), "Order", -demand])
            sheet.append([material, "DEMO-PLANT", monday + timedelta(days=4), "PurRqs", supply])

    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="17394A")
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 28
    for column, width in {"A": 25, "B": 20, "C": 49, "D": 28, "E": 14}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "@"
        row[1].number_format = "@"
        row[2].number_format = "yyyy-mm-dd"
        row[4].number_format = "0"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output)
    print(f"Created {output.name}: {sheet.max_row - 1} records, 3 materials, 32 weeks.")


if __name__ == "__main__":
    main()
